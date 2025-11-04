"""
授权书生成系统 - 上传生成文件模块

提供授权书数据管理和文档生成功能（Excel导入、数据CRUD、PDF生成）
详细接口文档请查看：UPLOAD_GENERATED_API.md
"""

from flask import Blueprint, render_template, request, jsonify, send_file
from app.utils.decorators import admin_required, generated_required, generate_admin_required
from app.models.authorization import Authorization
from app.models.generated import Generated
from app.models.user import User
from app.extensions import db
from app.utils.pdf_generator import generate_pdf_document
from app.utils.task_manager import task_manager
from datetime import datetime, timedelta
import os
import re
import zipfile
import threading
from io import BytesIO

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


upload_generated_bp = Blueprint('upload_generated', __name__)

# Excel表头期望值（用于数据导入时的格式校验）
EXPECTED_HEADERS = [
    '使用模板', '店铺类型', '授权主体', '授权平台', '授权品牌', '品牌商标号',
    '店铺名称', '店铺ID', '所属分公司', '授权期间', '授权字号', '授权方（盖章）主体', '用印时间'
]

"""功能菜单页面"""
@upload_generated_bp.route('/upload-generated/menu')
@generate_admin_required
def menu_page():
    return render_template('uploadGeneratedFiles/menu.html')

"""Excel数据导入页面"""
@upload_generated_bp.route('/upload-generated/import')
@generate_admin_required
def import_page():
    return render_template('uploadGeneratedFiles/upload.html')

"""下载Excel导入模板"""
@upload_generated_bp.route('/upload-generated/download-template')
@generate_admin_required
def download_template():
    """下载Excel导入模板文件"""
    from flask import current_app
    # 使用Flask应用的根路径构建绝对路径
    template_path = os.path.join(current_app.root_path, 'static', 'templates', '授权书上传模板.xlsx')
    
    if not os.path.exists(template_path):
        return jsonify({'success': False, 'message': '模板文件不存在'}), 404
    
    return send_file(
        template_path,
        as_attachment=True,
        download_name='授权书上传模板.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

"""数据列表查看页面"""
@upload_generated_bp.route('/upload-generated/list')
@generated_required
def list_page():
    return render_template('uploadGeneratedFiles/list.html')

"""获取数据列表（支持分页、筛选、模糊搜索）"""
@upload_generated_bp.route('/upload-generated/api/list', methods=['GET'])
@generated_required
def api_list():
    try:
        from flask import session
        
        # 获取分页参数
        page = request.args.get('page', 1, type=int)
        page_size = request.args.get('page_size', 10, type=int)
        
        # 获取筛选参数
        template_used = request.args.get('template_used', '').strip()
        store_type = request.args.get('store_type', '').strip()
        authorized_entity = request.args.get('authorized_entity', '').strip()
        platform = request.args.get('platform', '').strip()
        brand = request.args.get('brand', '').strip()
        store_name = request.args.get('store_name', '').strip()
        auth_number = request.args.get('auth_number', '').strip()
        stamping_entity = request.args.get('stamping_entity', '').strip()
        
        # 构建查询
        query = Generated.query
        
        # 根据用户角色过滤数据
        # generated角色只能看到所属分公司一致的数据
        # generate_admin角色可以看到所有数据
        user_role = session.get('role')
        user_entity = session.get('entity')
        
        if user_role == 'generated' and user_entity:
            # 普通用户只能看到所属分公司匹配的数据
            query = query.filter(Generated.subsidiary == user_entity)
        # generate_admin不需要额外过滤，可以看到所有数据
        
        # 应用筛选条件
        if template_used:
            query = query.filter(Generated.template_used.like(f'%{template_used}%'))
        if store_type:
            query = query.filter(Generated.store_type.like(f'%{store_type}%'))
        if authorized_entity:
            query = query.filter(Generated.authorized_entity.like(f'%{authorized_entity}%'))
        if platform:
            query = query.filter(Generated.platform.like(f'%{platform}%'))
        if brand:
            query = query.filter(Generated.brand.like(f'%{brand}%'))
        if store_name:
            query = query.filter(Generated.store_name.like(f'%{store_name}%'))
        if auth_number:
            query = query.filter(Generated.auth_number.like(f'%{auth_number}%'))
        if stamping_entity:
            query = query.filter(Generated.stamping_entity.like(f'%{stamping_entity}%'))
        
        # 获取总数
        total = query.count()
        
        # 分页查询
        query = query.order_by(Generated.created_at.desc())
        pagination = query.paginate(page=page, per_page=page_size, error_out=False)
        
        # 转换为字典
        items = [item.to_dict() for item in pagination.items]
        
        return jsonify({
            'success': True,
            'data': items,
            'total': total,
            'page': page,
            'page_size': page_size,
            'total_pages': pagination.pages
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'查询失败: {str(e)}'}), 500

"""获取筛选器选项（所有字段的唯一值列表）"""
@upload_generated_bp.route('/upload-generated/api/filters', methods=['GET'])
@generated_required
def api_filters():
    try:
        from flask import session
        
        # 根据用户角色过滤数据
        user_role = session.get('role')
        user_entity = session.get('entity')
        
        # 构建基础查询
        base_query = db.session.query(Generated)
        
        # 普通用户只能看到自己分公司的数据
        if user_role == 'generated' and user_entity:
            base_query = base_query.filter(Generated.subsidiary == user_entity)
        
        # 获取所有唯一值用于下拉选项
        templates = base_query.with_entities(Generated.template_used).distinct().filter(Generated.template_used != '').all()
        store_types = base_query.with_entities(Generated.store_type).distinct().filter(Generated.store_type != '').all()
        entities = base_query.with_entities(Generated.authorized_entity).distinct().filter(Generated.authorized_entity != '').all()
        platforms = base_query.with_entities(Generated.platform).distinct().filter(Generated.platform != '').all()
        brands = base_query.with_entities(Generated.brand).distinct().filter(Generated.brand != '').all()
        stamping_entities = base_query.with_entities(Generated.stamping_entity).distinct().filter(Generated.stamping_entity != '').all()
        
        return jsonify({
            'success': True,
            'filters': {
                'templates': [t[0] for t in templates if t[0]],
                'store_types': [s[0] for s in store_types if s[0]],
                'entities': [e[0] for e in entities if e[0]],
                'platforms': [p[0] for p in platforms if p[0]],
                'brands': [b[0] for b in brands if b[0]],
                'stamping_entities': [s[0] for s in stamping_entities if s[0]]
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取筛选器失败: {str(e)}'}), 500

"""获取单条记录详情"""
@upload_generated_bp.route('/upload-generated/api/record/<int:id>', methods=['GET'])
@generated_required
def api_get_record(id):
    
    try:
        record = Generated.query.get(id)
        if not record:
            return jsonify({'success': False, 'message': '记录不存在'}), 404
        
        return jsonify({
            'success': True,
            'data': record.to_dict()
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取记录失败: {str(e)}'}), 500

"""更新记录（支持事务回滚）"""
@upload_generated_bp.route('/upload-generated/api/record/<int:id>', methods=['PUT'])
@generate_admin_required
def api_update_record(id):
    
    try:
        record = Generated.query.get(id)
        if not record:
            return jsonify({'success': False, 'message': '记录不存在'}), 404
        
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'message': '无效的数据'}), 400
        
        # 更新字段
        record.template_used = data.get('template_used', '').strip()
        record.store_type = data.get('store_type', '').strip()
        record.authorized_entity = data.get('authorized_entity', '').strip()
        record.platform = data.get('platform', '').strip()
        record.brand = data.get('brand', '').strip()
        record.trademark_no = data.get('trademark_no', '').strip()
        record.store_name = data.get('store_name', '').strip()
        record.store_id = data.get('store_id', '').strip()
        record.subsidiary = data.get('subsidiary', '').strip()
        record.period = data.get('period', '').strip()
        record.auth_number = data.get('auth_number', '').strip()
        record.stamping_entity = data.get('stamping_entity', '').strip()
        record.stamping_date = data.get('stamping_date', '').strip()
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': '记录更新成功',
            'data': record.to_dict()
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'更新失败: {str(e)}'}), 500

"""删除记录（不可恢复）"""
@upload_generated_bp.route('/upload-generated/api/record/<int:id>', methods=['DELETE'])
@generate_admin_required
def api_delete_record(id):   
    try:
        record = Generated.query.get(id)
        if not record:
            return jsonify({'success': False, 'message': '记录不存在'}), 404
        
        db.session.delete(record)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': '记录删除成功'
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'删除失败: {str(e)}'}), 500


"""查询任务进度"""
@upload_generated_bp.route('/upload-generated/api/task/<task_id>', methods=['GET'])
@generated_required
def api_get_task_status(task_id):
    try:
        task_info = task_manager.get_task(task_id)
        
        if not task_info:
            return jsonify({'success': False, 'message': '任务不存在'}), 404
        
        return jsonify({
            'success': True,
            'task': task_info
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'查询失败: {str(e)}'}), 500


"""异步批量生成PDF文档"""
@upload_generated_bp.route('/upload-generated/api/generate-pdf-batch-async', methods=['POST'])
@generated_required
def api_generate_pdf_batch_async():
    try:
        from flask import current_app
        # 获取实际的应用实例（而不是代理对象）
        app = current_app._get_current_object()
        
        payload = request.get_json(silent=True) or {}
        ids = payload.get('ids', [])
        
        if not ids:
            return jsonify({'success': False, 'message': '请选择要生成的数据'}), 400
        
        # 查询数据
        records = Generated.query.filter(Generated.id.in_(ids)).all()
        
        if not records:
            return jsonify({'success': False, 'message': '未找到指定的数据'}), 404
        
        # 创建任务
        task_id = task_manager.create_task('pdf_batch', total=len(records))
        
        # 启动后台线程
        def generate_pdf_batch_task():
            # 在后台线程中使用应用上下文
            with app.app_context():
                try:
                    from app.utils.pdf_generator import PDFGenerator
                    generator = PDFGenerator()
                    
                    task_manager.set_task_running(task_id, f'正在批量生成 {len(records)} 个PDF文档...')
                    
                    success_count = 0
                    failed_count = 0
                    files = []
                    errors = []
                    
                    for index, record in enumerate(records, 1):
                        try:
                            task_manager.set_task_progress(
                                task_id,
                                index,
                                f'正在生成 {index}/{len(records)}: {record.store_name or "未命名"}'
                            )
                            
                            success, result = generator.generate_from_model(record)
                            if success:
                                success_count += 1
                                files.append(os.path.basename(result))
                            else:
                                failed_count += 1
                                errors.append(f'ID {record.id}: {result}')
                        except Exception as e:
                            failed_count += 1
                            errors.append(f'ID {record.id}: {str(e)}')
                    
                    task_manager.set_task_completed(
                        task_id,
                        result={
                            'success': success_count,
                            'failed': failed_count,
                            'files': files,
                            'errors': errors
                        },
                        message=f'批量生成完成：成功 {success_count} 个，失败 {failed_count} 个'
                    )
                except Exception as e:
                    task_manager.set_task_failed(task_id, str(e))
        
        thread = threading.Thread(target=generate_pdf_batch_task)
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'success': True,
            'task_id': task_id,
            'message': f'任务已创建，正在后台生成 {len(records)} 个PDF...'
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'创建任务失败: {str(e)}'}), 500


"""批量生成PDF文档"""
@upload_generated_bp.route('/upload-generated/api/generate-pdf-batch', methods=['POST'])
@generated_required
def api_generate_pdf_batch():
    try:
        payload = request.get_json(silent=True) or {}
        ids = payload.get('ids', [])
        
        if not ids:
            return jsonify({'success': False, 'message': '请选择要生成的数据'}), 400
        
        # 查询数据
        records = Generated.query.filter(Generated.id.in_(ids)).all()
        
        if not records:
            return jsonify({'success': False, 'message': '未找到指定的数据'}), 404
        
        # 批量生成PDF
        from app.utils.pdf_generator import PDFGenerator
        generator = PDFGenerator()
        
        success_count = 0
        failed_count = 0
        files = []
        errors = []
        
        for record in records:
            try:
                success, result = generator.generate_from_model(record)
                if success:
                    success_count += 1
                    files.append(os.path.basename(result))
                else:
                    failed_count += 1
                    errors.append(f'ID {record.id}: {result}')
            except Exception as e:
                failed_count += 1
                errors.append(f'ID {record.id}: {str(e)}')
        
        return jsonify({
            'success': True,
            'message': f'成功生成 {success_count} 个PDF，失败 {failed_count} 个',
            'results': {
                'success': success_count,
                'failed': failed_count,
                'files': files,
                'errors': errors
            }
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'批量生成PDF失败: {str(e)}'}), 500


"""异步生成PDF文档"""
@upload_generated_bp.route('/upload-generated/api/generate-pdf-async/<int:id>', methods=['POST'])
@generated_required
def api_generate_pdf_async(id):
    try:
        from flask import current_app
        # 获取实际的应用实例（而不是代理对象）
        app = current_app._get_current_object()
        
        # 查询数据
        record = Generated.query.get(id)
        if not record:
            return jsonify({'success': False, 'message': '数据不存在'}), 404
        
        # 创建任务
        task_id = task_manager.create_task('pdf', total=1)
        
        # 启动后台线程
        def generate_pdf_task():
            # 在后台线程中使用应用上下文
            with app.app_context():
                try:
                    task_manager.set_task_running(task_id, '正在生成PDF文档...')
                    # 模拟进度：开始
                    task_manager.update_task(task_id, progress=10, current=0, message='正在读取模板...')
                    
                    # 模拟进度：准备中
                    import time
                    time.sleep(0.2)
                    task_manager.update_task(task_id, progress=25, current=0, message='正在填充数据...')
                    
                    time.sleep(0.2)
                    task_manager.update_task(task_id, progress=40, current=0, message='正在生成PDF文档...')
                    
                    success, result = generate_pdf_document(record)
                    
                    if success:
                        # 模拟进度：转换中
                        task_manager.update_task(task_id, progress=80, current=0, message='正在转换为PDF...')
                        time.sleep(0.3)
                        task_manager.update_task(task_id, progress=95, current=0, message='正在保存文件...')
                        time.sleep(0.1)
                        
                        task_manager.set_task_completed(
                            task_id,
                            result={
                                'filename': os.path.basename(result),
                                'file_path': result.replace('\\', '/').replace('app/', '')
                            },
                            message='PDF文档生成成功'
                        )
                    else:
                        task_manager.set_task_failed(task_id, result)
                except Exception as e:
                    task_manager.set_task_failed(task_id, str(e))
        
        thread = threading.Thread(target=generate_pdf_task)
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'success': True,
            'task_id': task_id,
            'message': '任务已创建，正在后台生成...'
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'创建任务失败: {str(e)}'}), 500


"""生成PDF文档"""
@upload_generated_bp.route('/upload-generated/api/generate-pdf/<int:id>', methods=['POST'])
@generated_required
def api_generate_pdf(id):
    
    try:
        # 查询数据
        record = Generated.query.get(id)
        if not record:
            return jsonify({'success': False, 'message': '数据不存在'}), 404
        
        # 生成PDF文档
        success, result = generate_pdf_document(record)
        
        if success:
            # 返回文件路径（相对路径）
            relative_path = result.replace('\\', '/').replace('app/', '')
            return jsonify({
                'success': True,
                'message': 'PDF生成成功',
                'file_path': relative_path,
                'filename': os.path.basename(result)
            })
        else:
            return jsonify({'success': False, 'message': result}), 500
            
    except Exception as e:
        return jsonify({'success': False, 'message': f'生成PDF失败: {str(e)}'}), 500

"""检查PDF文件是否已存在"""
@upload_generated_bp.route('/upload-generated/api/check-pdf/<int:id>', methods=['GET'])
@generated_required
def api_check_pdf(id):
    try:
        from flask import current_app
        
        # 查找记录
        record = Generated.query.get(id)
        if not record:
            return jsonify({'success': False, 'message': '记录不存在'}), 404
        
        # 生成预期的文件名（与pdf_generator中的逻辑一致）
        parts = []
        if record.store_type:
            parts.append(record.store_type)
        if record.store_name:
            parts.append(record.store_name)
        
        filename = '_'.join(parts)
        # 清理文件名中的非法字符
        import re
        filename = re.sub(r'[<>:"/\\|?*]', '_', filename)
        filename = f"{filename}.pdf"
        
        # 检查文件是否存在
        output_dir = current_app.config.get('UPLOAD_FOLDER_SHOUQUAN', 'app/static/generated_docs')
        file_path = os.path.join(output_dir, filename)
        
        exists = os.path.exists(file_path)
        
        return jsonify({
            'success': True,
            'exists': exists,
            'filename': filename if exists else None
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'检查失败: {str(e)}'}), 500


@upload_generated_bp.route('/upload-generated/preview-pdf/<path:filename>')
@generated_required
def preview_pdf(filename):
    """在浏览器中预览PDF文档"""
    try:
        from flask import current_app
        # 使用配置中的路径
        output_dir = current_app.config.get('UPLOAD_FOLDER_SHOUQUAN', 'app/static/generated_docs')
        file_path = os.path.join(output_dir, filename)
        
        if not os.path.exists(file_path):
            return jsonify({'success': False, 'message': 'PDF文件不存在'}), 404
        
        return send_file(
            file_path,
            mimetype='application/pdf',
            as_attachment=False,  # 在浏览器中预览而不是下载
            download_name=filename
        )
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'预览失败: {str(e)}'}), 500


@upload_generated_bp.route('/upload-generated/download-pdf/<path:filename>')
@generated_required
def download_pdf(filename):
    """下载PDF文档"""
    try:
        from flask import current_app
        # 使用配置中的路径
        output_dir = current_app.config.get('UPLOAD_FOLDER_SHOUQUAN', 'app/static/generated_docs')
        file_path = os.path.join(output_dir, filename)
        
        if not os.path.exists(file_path):
            return jsonify({'success': False, 'message': 'PDF文件不存在'}), 404
        
        return send_file(
            file_path,
            as_attachment=True,
            download_name=filename,
            mimetype='application/pdf'
        )
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'下载PDF失败: {str(e)}'}), 500


"""批量下载文档（打包成ZIP，自动生成缺失文件）"""
@upload_generated_bp.route('/upload-generated/api/download-batch', methods=['POST'])
@generated_required
def api_download_batch():
    try:
        from flask import current_app
        
        payload = request.get_json(silent=True) or {}
        ids = payload.get('ids', [])
        file_type = payload.get('type', 'pdf')  # 只支持 'pdf'
        
        if not ids:
            return jsonify({'success': False, 'message': '请选择要下载的文件'}), 400
        
        # 查询数据
        records = Generated.query.filter(Generated.id.in_(ids)).all()
        
        if not records:
            return jsonify({'success': False, 'message': '未找到指定的数据'}), 404
        
        # 使用配置中的路径
        output_dir = current_app.config.get('UPLOAD_FOLDER_SHOUQUAN', 'app/static/generated_docs')
        
        # 统计信息
        generated_count = 0
        existed_count = 0
        failed_count = 0
        failed_files = []
        
        # 创建ZIP文件
        memory_file = BytesIO()
        
        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            for record in records:
                # 生成文件名
                parts = []
                if record.store_type:
                    parts.append(record.store_type)
                if record.store_name:
                    parts.append(record.store_name)
                
                if not parts:
                    failed_count += 1
                    failed_files.append(f'ID {record.id}: 缺少店铺类型或店铺名称')
                    continue
                
                filename = '_'.join(parts)
                # 清理文件名中的非法字符
                filename = re.sub(r'[<>:"/\\|?*]', '_', filename)
                
                # 只支持PDF
                filename = f"{filename}.pdf"
                
                file_path = os.path.join(output_dir, filename)
                
                # 检查文件是否存在
                if not os.path.exists(file_path):
                    # 文件不存在，先生成
                    try:
                        # 生成PDF
                        success, result = generate_pdf_document(record)
                        
                        if success:
                            generated_count += 1
                            # 重新获取文件路径（因为生成函数返回的是完整路径）
                            file_path = result
                        else:
                            failed_count += 1
                            failed_files.append(f'{filename}: {result}')
                            continue
                    except Exception as e:
                        failed_count += 1
                        failed_files.append(f'{filename}: {str(e)}')
                        continue
                else:
                    existed_count += 1
                
                # 添加文件到ZIP
                if os.path.exists(file_path):
                    zf.write(file_path, filename)
        
        total_files = generated_count + existed_count
        
        if total_files == 0:
            return jsonify({
                'success': False, 
                'message': f'所有文件生成失败，无法下载。失败原因：{"; ".join(failed_files[:3])}'
            }), 500
        
        # 重置文件指针到开始位置
        memory_file.seek(0)
        
        # 生成ZIP文件名
        timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
        zip_filename = f'授权书_{file_type}_{timestamp}.zip'
        
        # 在响应头中添加统计信息
        response = send_file(
            memory_file,
            mimetype='application/zip',
            as_attachment=True,
            download_name=zip_filename
        )
        
        # 添加自定义响应头传递统计信息
        response.headers['X-Generated-Count'] = str(generated_count)
        response.headers['X-Existed-Count'] = str(existed_count)
        response.headers['X-Failed-Count'] = str(failed_count)
        response.headers['X-Total-Count'] = str(total_files)
        
        return response
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'批量下载失败: {str(e)}'}), 500


"""异步批量下载文档（打包成ZIP，自动生成缺失文件，显示进度）"""
@upload_generated_bp.route('/upload-generated/api/download-batch-async', methods=['POST'])
@generated_required
def api_download_batch_async():
    try:
        from flask import current_app
        # 获取实际的应用实例
        app = current_app._get_current_object()
        
        payload = request.get_json(silent=True) or {}
        ids = payload.get('ids', [])
        file_type = payload.get('type', 'pdf')  # 只支持 'pdf'
        
        if not ids:
            return jsonify({'success': False, 'message': '请选择要下载的文件'}), 400
        
        # 查询数据
        records = Generated.query.filter(Generated.id.in_(ids)).all()
        
        if not records:
            return jsonify({'success': False, 'message': '未找到指定的数据'}), 404
        
        # 创建任务
        task_id = task_manager.create_task('download_batch', total=len(records))
        
        # 启动后台线程
        def download_batch_task():
            # 在后台线程中使用应用上下文
            with app.app_context():
                try:
                    task_manager.set_task_running(task_id, f'正在准备批量下载 {len(records)} 个{file_type.upper()}文档...')
                    
                    output_dir = app.config.get('UPLOAD_FOLDER_SHOUQUAN', 'app/static/generated_docs')
                    
                    # 统计信息
                    generated_count = 0
                    existed_count = 0
                    failed_count = 0
                    failed_files = []
                    
                    # 创建ZIP文件
                    memory_file = BytesIO()
                    
                    with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
                        for index, record in enumerate(records, 1):
                            try:
                                task_manager.set_task_progress(
                                    task_id,
                                    index,
                                    f'正在处理 {index}/{len(records)}: {record.store_name or "未命名"}'
                                )
                                
                                # 生成文件名
                                parts = []
                                if record.store_type:
                                    parts.append(record.store_type)
                                if record.store_name:
                                    parts.append(record.store_name)
                                
                                if not parts:
                                    failed_count += 1
                                    failed_files.append(f'ID {record.id}: 缺少店铺类型或店铺名称')
                                    continue
                                
                                filename = '_'.join(parts)
                                filename = re.sub(r'[<>:"/\\|?*]', '_', filename)
                                
                                # 只支持PDF
                                filename = f"{filename}.pdf"
                                
                                file_path = os.path.join(output_dir, filename)
                                
                                # 检查文件是否存在
                                if not os.path.exists(file_path):
                                    # 文件不存在，先生成
                                    try:
                                        success, result = generate_pdf_document(record)
                                        
                                        if success:
                                            generated_count += 1
                                            file_path = result
                                        else:
                                            failed_count += 1
                                            failed_files.append(f'{filename}: {result}')
                                            continue
                                    except Exception as e:
                                        failed_count += 1
                                        failed_files.append(f'{filename}: {str(e)}')
                                        continue
                                else:
                                    existed_count += 1
                                
                                # 添加文件到ZIP
                                if os.path.exists(file_path):
                                    zf.write(file_path, filename)
                            except Exception as e:
                                failed_count += 1
                                failed_files.append(f'ID {record.id}: {str(e)}')
                    
                    total_files = generated_count + existed_count
                    
                    if total_files == 0:
                        task_manager.set_task_failed(
                            task_id, 
                            f'所有文件生成失败。失败原因：{"; ".join(failed_files[:3])}'
                        )
                        return
                    
                    # 保存ZIP文件到临时位置
                    memory_file.seek(0)
                    timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
                    zip_filename = f'授权书_{file_type}_{timestamp}.zip'
                    temp_zip_path = os.path.join(output_dir, f'temp_{zip_filename}')
                    
                    with open(temp_zip_path, 'wb') as f:
                        f.write(memory_file.getvalue())
                    
                    task_manager.set_task_completed(
                        task_id,
                        result={
                            'zip_filename': zip_filename,
                            'zip_path': temp_zip_path,
                            'generated_count': generated_count,
                            'existed_count': existed_count,
                            'failed_count': failed_count,
                            'total_count': total_files
                        },
                        message=f'批量下载准备完成：成功 {total_files} 个，失败 {failed_count} 个'
                    )
                    
                except Exception as e:
                    task_manager.set_task_failed(task_id, str(e))
        
        thread = threading.Thread(target=download_batch_task)
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'success': True,
            'task_id': task_id,
            'message': '任务已创建，正在后台准备下载...'
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'创建下载任务失败: {str(e)}'}), 500


"""下载异步批量准备好的ZIP文件"""
@upload_generated_bp.route('/upload-generated/api/download-zip/<filename>', methods=['GET'])
@generated_required
def api_download_zip(filename):
    try:
        from flask import current_app
        
        output_dir = current_app.config.get('UPLOAD_FOLDER_SHOUQUAN', 'app/static/generated_docs')
        temp_zip_path = os.path.join(output_dir, f'temp_{filename}')
        
        if not os.path.exists(temp_zip_path):
            return jsonify({'success': False, 'message': 'ZIP文件不存在或已过期'}), 404
        
        response = send_file(
            temp_zip_path,
            mimetype='application/zip',
            as_attachment=True,
            download_name=filename
        )
        
        # 下载完成后删除临时文件（延迟删除）
        @response.call_on_close
        def cleanup():
            try:
                import time
                time.sleep(2)  # 等待2秒确保下载完成
                if os.path.exists(temp_zip_path):
                    os.remove(temp_zip_path)
            except:
                pass
        
        return response
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'下载失败: {str(e)}'}), 500


@upload_generated_bp.route('/upload-generated/preview', methods=['POST'])
@generate_admin_required
def preview_excel():
    """预览Excel数据（校验表头，返回数据行）"""
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '未选择文件'})
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': '未选择文件名'})
    err, rows = _parse_excel(file)
    if err:
        return jsonify({'success': False, 'message': err.get('error'), 'expected': err.get('expected'), 'found': err.get('found')}), 400
    return jsonify({'success': True, 'rows': rows, 'count': len(rows)})


@upload_generated_bp.route('/upload-generated/confirm', methods=['POST'])
@generate_admin_required
def confirm_import():
    """确认导入Excel数据到数据库（支持事务回滚）"""
    payload = request.get_json(silent=True) or {}
    rows = payload.get('rows', [])
    if not isinstance(rows, list) or not rows:
        return jsonify({'success': False, 'message': '无可导入数据'})

    inserted = 0
    try:
        for item in rows:
            record = Generated(
                template_used=(item.get('使用模板') or '').strip(),
                store_type=(item.get('店铺类型') or '').strip(),
                authorized_entity=(item.get('授权主体') or '').strip(),
                platform=(item.get('授权平台') or '').strip(),
                brand=(item.get('授权品牌') or '').strip(),
                trademark_no=(item.get('品牌商标号') or '').strip(),
                store_name=(item.get('店铺名称') or '').strip(),
                store_id=(item.get('店铺ID') or '').strip(),
                subsidiary=(item.get('所属分公司') or '').strip(),
                period=(item.get('授权期间') or '').strip(),
                auth_number=(item.get('授权字号') or '').strip(),
                stamping_entity=(item.get('授权方（盖章）主体') or '').strip(),
                stamping_date=(item.get('用印时间') or '').strip(),
            )
            db.session.add(record)
            inserted += 1
        db.session.commit()
        return jsonify({'success': True, 'inserted': inserted})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'导入失败: {str(e)}'}), 500


def _format_excel_value(value):
    """格式化Excel单元格值，特别处理日期格式"""
    if value is None:
        return ''
    
    # 如果是datetime对象，格式化为中文日期
    if isinstance(value, datetime):
        return f"{value.year}年{value.month}月{value.day}日"
    
    # 如果是字符串，直接返回
    if isinstance(value, str):
        return value.strip()
    
    # 如果是数字，可能是Excel的日期序列号
    # Excel日期从1900年1月1日开始，序列号在25000-60000之间通常是日期
    if isinstance(value, (int, float)) and 25000 <= value <= 60000:
        try:
            # Excel的日期基准是1899年12月30日（因为1900年bug）
            base_date = datetime(1899, 12, 30)
            date_value = base_date + timedelta(days=value)
            return f"{date_value.year}年{date_value.month}月{date_value.day}日"
        except:
            # 如果转换失败，返回原始值的字符串
            return str(value).strip()
    
    # 其他类型转为字符串
    return str(value).strip()


def _parse_excel(file_storage):
    """
    内部函数：解析Excel文件（使用openpyxl）
    
    特性：
    - 支持任意列顺序，根据表头名称智能映射
    - 缺少的列会自动填充空值
    - 多余的列会被忽略
    - 表头名称必须完全匹配（区分大小写）
    """
    if load_workbook is None:
        return {'error': '服务器缺少 openpyxl 依赖，请安装后重试'}, None
    try:
        wb = load_workbook(filename=file_storage, data_only=True)
        ws = wb.active
        
        # 读取表头
        headers = [
            (cell.value or '').strip() if isinstance(cell.value, str) else str(cell.value or '')
            for cell in next(ws.iter_rows(min_row=1, max_row=1))
        ]
        
        # 检查是否包含至少一个期望的表头
        valid_headers = [h for h in headers if h in EXPECTED_HEADERS]
        if not valid_headers:
            return {
                'error': '未找到有效的表头列',
                'expected': EXPECTED_HEADERS,
                'found': headers
            }, None
        
        # 创建表头索引映射 {表头名: 列索引}
        header_index_map = {header: idx for idx, header in enumerate(headers)}
        
        # 读取数据行
        data_rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            # 根据表头映射构建记录
            record = {}
            for expected_header in EXPECTED_HEADERS:
                if expected_header in header_index_map:
                    col_idx = header_index_map[expected_header]
                    # 确保索引不越界
                    if col_idx < len(row):
                        value = row[col_idx]
                        record[expected_header] = _format_excel_value(value)
                    else:
                        record[expected_header] = ''
                else:
                    # Excel中没有这一列，填充空值
                    record[expected_header] = ''
            
            # 跳过完全空行
            if any(v for v in record.values()):
                data_rows.append(record)
        
        return None, data_rows
    except Exception as e:
        return {'error': f'解析失败: {str(e)}'}, None


# ==================== 用户管理相关路由 ====================

"""用户管理页面"""
@upload_generated_bp.route('/upload-generated/user-management')
@generate_admin_required
def user_management_page():
    return render_template('uploadGeneratedFiles/user_management.html')

"""获取generated角色用户列表"""
@upload_generated_bp.route('/upload-generated/api/users', methods=['GET'])
@generate_admin_required
def api_get_users():
    try:
        # 获取分页参数
        page = request.args.get('page', 1, type=int)
        page_size = request.args.get('page_size', 10, type=int)
        
        # 获取搜索参数
        search = request.args.get('search', '').strip()
        
        # 构建查询 - 只查询generated和generate_admin角色的用户
        query = User.query.filter(User.role.in_(['generated', 'generate_admin']))
        
        # 应用搜索条件
        if search:
            query = query.filter(
                db.or_(
                    User.username.like(f'%{search}%'),
                    User.entity.like(f'%{search}%')
                )
            )
        
        # 获取总数
        total = query.count()
        
        # 分页查询 - 管理员优先，然后按ID倒序
        # 使用 case 语句：管理员(generate_admin)排序为0，普通用户为1，然后按ID倒序
        from sqlalchemy import case
        query = query.order_by(
            case(
                (User.role == 'generate_admin', 0),
                else_=1
            ),
            User.id.desc()
        )
        pagination = query.paginate(page=page, per_page=page_size, error_out=False)
        
        # 转换为字典
        items = [{
            'id': user.id,
            'username': user.username,
            'role': user.role,
            'entity': user.entity or ''
        } for user in pagination.items]
        
        return jsonify({
            'success': True,
            'data': items,
            'total': total,
            'page': page,
            'page_size': page_size,
            'total_pages': pagination.pages
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'查询失败: {str(e)}'}), 500

"""创建新用户"""
@upload_generated_bp.route('/upload-generated/api/users', methods=['POST'])
@generate_admin_required
def api_create_user():
    try:
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'message': '无效的数据'}), 400
        
        username = data.get('username', '').strip()
        password = data.get('password', '').strip()
        role = data.get('role', '').strip()
        entity = data.get('entity', '').strip()
        
        # 验证必填字段
        if not username:
            return jsonify({'success': False, 'message': '用户名不能为空'}), 400
        if not password:
            return jsonify({'success': False, 'message': '密码不能为空'}), 400
        if not entity:
            return jsonify({'success': False, 'message': '所属分公司不能为空'}), 400
        
        # 新增用户时只能创建普通用户（generated角色）
        if role and role != 'generated':
            return jsonify({'success': False, 'message': '新增用户只能创建普通用户角色'}), 400
        
        # 强制设置为generated角色
        role = 'generated'
        
        # 检查用户名是否已存在
        existing_user = User.query.filter_by(username=username).first()
        if existing_user:
            return jsonify({'success': False, 'message': '用户名已存在'}), 400
        
        # 创建新用户
        new_user = User(
            username=username,
            password=password,
            role=role,
            entity=entity
        )
        
        db.session.add(new_user)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': '用户创建成功',
            'data': {
                'id': new_user.id,
                'username': new_user.username,
                'role': new_user.role,
                'entity': new_user.entity or ''
            }
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'创建失败: {str(e)}'}), 500

"""获取单个用户信息"""
@upload_generated_bp.route('/upload-generated/api/users/<int:user_id>', methods=['GET'])
@generate_admin_required
def api_get_user(user_id):
    try:
        user = User.query.get(user_id)
        if not user:
            return jsonify({'success': False, 'message': '用户不存在'}), 404
        
        # 只允许查看generated相关角色的用户
        if user.role not in ['generated', 'generate_admin']:
            return jsonify({'success': False, 'message': '无权访问此用户'}), 403
        
        return jsonify({
            'success': True,
            'data': {
                'id': user.id,
                'username': user.username,
                'role': user.role,
                'entity': user.entity or ''
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取用户失败: {str(e)}'}), 500

"""更新用户信息"""
@upload_generated_bp.route('/upload-generated/api/users/<int:user_id>', methods=['PUT'])
@generate_admin_required
def api_update_user(user_id):
    try:
        user = User.query.get(user_id)
        if not user:
            return jsonify({'success': False, 'message': '用户不存在'}), 404
        
        # 只允许修改generated相关角色的用户
        if user.role not in ['generated', 'generate_admin']:
            return jsonify({'success': False, 'message': '无权修改此用户'}), 403
        
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'message': '无效的数据'}), 400
        
        # 如果是管理员账户，只能修改密码，不能修改其他信息
        if user.role == 'generate_admin':
            # 只处理密码更新
            if 'password' in data and data['password'].strip():
                user.set_password(data['password'].strip())
                db.session.commit()
                return jsonify({
                    'success': True,
                    'message': '密码更新成功',
                    'data': {
                        'id': user.id,
                        'username': user.username,
                        'role': user.role,
                        'entity': user.entity or ''
                    }
                })
            else:
                return jsonify({'success': False, 'message': '管理员账户只能修改密码'}), 400
        
        # 普通用户可以修改所有信息
        # 更新用户名
        if 'username' in data:
            new_username = data['username'].strip()
            if not new_username:
                return jsonify({'success': False, 'message': '用户名不能为空'}), 400
            # 检查用户名是否被其他用户使用
            existing_user = User.query.filter(User.username == new_username, User.id != user_id).first()
            if existing_user:
                return jsonify({'success': False, 'message': '用户名已被使用'}), 400
            user.username = new_username
        
        # 更新密码（如果提供）
        if 'password' in data and data['password'].strip():
            user.set_password(data['password'].strip())
        
        # 更新角色（普通用户不能升级为管理员）
        if 'role' in data:
            new_role = data['role'].strip()
            if new_role == 'generate_admin':
                return jsonify({'success': False, 'message': '不能将普通用户升级为管理员'}), 400
            if new_role not in ['generated']:
                return jsonify({'success': False, 'message': '角色只能是generated'}), 400
            user.role = new_role
        
        # 更新所属分公司（必填）
        if 'entity' in data:
            new_entity = data['entity'].strip()
            if not new_entity:
                return jsonify({'success': False, 'message': '所属分公司不能为空'}), 400
            user.entity = new_entity
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': '用户更新成功',
            'data': {
                'id': user.id,
                'username': user.username,
                'role': user.role,
                'entity': user.entity or ''
            }
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'更新失败: {str(e)}'}), 500

"""删除用户"""
@upload_generated_bp.route('/upload-generated/api/users/<int:user_id>', methods=['DELETE'])
@generate_admin_required
def api_delete_user(user_id):
    try:
        user = User.query.get(user_id)
        if not user:
            return jsonify({'success': False, 'message': '用户不存在'}), 404
        
        # 只允许删除generated相关角色的用户
        if user.role not in ['generated', 'generate_admin']:
            return jsonify({'success': False, 'message': '无权删除此用户'}), 403
        
        # 不允许删除管理员账户
        if user.role == 'generate_admin':
            return jsonify({'success': False, 'message': '管理员账户不允许删除'}), 403
        
        username = user.username
        db.session.delete(user)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'用户 {username} 删除成功'
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'删除失败: {str(e)}'}), 500


